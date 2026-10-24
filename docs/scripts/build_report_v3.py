"""Revise the user's manually edited v2 without rebuilding its cover or team data."""
from pathlib import Path
from copy import deepcopy
import hashlib
import json
import re
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.shared import Pt, RGBColor
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, Polygon

from scratch_report_builder_core import parse_body_text_segments, set_run_font

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'output/v2/Distributed_Design_Report_Draft_v2.docx'
WORK = ROOT / 'docs/v3-work'
STEM = '20261024-misttable-distributed-design-v3'
DOCX = WORK / f'{STEM}.docx'
WORK.mkdir(exist_ok=True)


def fill(p, text, size=10.5, refs=False):
    p.clear()
    for value, sup in ([(text, False)] if refs else parse_body_text_segments(text)):
        r = p.add_run(value)
        set_run_font(r, size_pt=size, superscript=sup)
    return p


def find(doc, prefix):
    found = [p for p in doc.paragraphs if p.text.startswith(prefix)]
    if len(found) != 1:
        raise ValueError(f'Expected one paragraph: {prefix!r}; found {len(found)}')
    return found[0]


def replace(doc, prefix, text):
    return fill(find(doc, prefix), text)


def before(doc, anchor_prefix, text):
    anchor = find(doc, anchor_prefix)
    sample = find(doc, '根据分布式系统经典定义')
    el = OxmlElement('w:p')
    el.append(deepcopy(sample._p.pPr))
    anchor._p.addprevious(el)
    p = Paragraph(el, anchor._parent)
    return fill(p, text)


def table(doc, number):
    cap = find(doc, f'表 {number}  ')
    el = cap._p.getnext()
    if el.tag != qn('w:tbl'):
        raise ValueError(f'Table {number} not adjacent to caption')
    return Table(el, doc)


def cell(t, row, col, text):
    c = t.cell(row, col)
    fill(c.paragraphs[0], text, size=9)
    for p in c.paragraphs[1:]:
        p._p.getparent().remove(p._p)


def draw_diagrams():
    out = WORK / 'diagrams'
    out.mkdir(exist_ok=True)
    plt.rcParams['font.sans-serif'] = ['Microsoft YaHei', 'SimHei']
    plt.rcParams['axes.unicode_minus'] = False

    def canvas():
        fig, ax = plt.subplots(figsize=(10, 6.2), dpi=260)
        ax.set(xlim=(0, 100), ylim=(0, 100))
        ax.axis('off')
        return fig, ax

    def box(ax, x, y, w, h, text, color='#e0f2fe', size=9):
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle='round,pad=0.6',
                                  facecolor=color, edgecolor='#64748b', lw=1.1))
        ax.text(x+w/2, y+h/2, text, ha='center', va='center', fontsize=size,
                color='#1e293b', linespacing=1.65)

    def arrow(ax, x1, y1, x2, y2):
        ax.annotate('', xy=(x2, y2), xytext=(x1, y1),
                    arrowprops={'arrowstyle': '->', 'color': '#475569', 'lw': 1.2})

    fig, ax = canvas()
    box(ax, 2, 39, 28, 53, '服务端权威状态\n\n预置地图原始文件\n全部角色与隐藏实体\n迷雾规则、权限版本\n房间持久化事件序号', '#f1f5f9')
    box(ax, 36, 43, 24, 45, '授权投影引擎\n\n角色、权限与迷雾\n地图瓦片服务端裁剪\n实体新增 / 更新 / 移除\n生成成员专属视图流', '#dcfce7', 8.6)
    box(ax, 67, 66, 31, 26, '主持人视图\n完整地图与全部实体\n控制迷雾及成员权限', '#fef3c7')
    box(ax, 67, 29, 31, 30, '玩家受限视图\n仅授权瓦片及可见实体\n隐藏坐标、属性不下发\n视图序号连续，隐藏事件不占号', '#e0f2fe', 8.4)
    arrow(ax, 30, 65, 36, 65)
    arrow(ax, 60, 72, 67, 78)
    arrow(ax, 60, 54, 67, 45)
    box(ax, 2, 2, 96, 20, '一致的访问边界：实时推送、重连快照与地图瓦片使用相同鉴权规则。\n对象由可见变为不可见时发出移除；权限版本变化时更换视图流并重置快照。\n原始底图不向玩家公开；已看见或下载的信息无法追溯撤回。', '#f8fafc', 8.8)
    fig.tight_layout()
    fig.savefig(out/'fig5_permission_projection.png', bbox_inches='tight')
    plt.close(fig)

    fig, ax = canvas()
    box(ax, 21, 85, 58, 11, '重新建连并提交身份凭证、视图流标识、最后确认的 viewSeq', size=8.6)
    box(ax, 21, 68, 58, 11, '重新鉴权；检查权限版本及数据历史代次\n存储回退时暂停恢复，先完成显式历史重置', '#fef3c7', 8.6)
    ax.add_patch(Polygon([[50, 62], [78, 50], [50, 38], [22, 50]],
                         facecolor='#fef08a', edgecolor='#a16207', lw=1.1))
    ax.text(50, 50, '同一视图流及权限版本，\n且可见流补发记录仍完整？', ha='center', va='center', fontsize=9)
    box(ax, 2, 14, 44, 19, '增量恢复\n补发连续的授权视图事件\n实体更新 / 移除按 viewSeq 应用\n不按房间序号跳号判断漏收', '#dcfce7', 8.5)
    box(ax, 54, 14, 44, 19, '授权快照重置\n固定基线 B 并缓冲 B 之后事件\n整体替换旧视图，更换流标识\n快照确认后接续缓冲事件', '#e0f2fe', 8.5)
    box(ax, 18, 1, 64, 8, '确认连续应用完成，恢复实时协作；缺口或版本变化重新同步', '#dcfce7', 8.5)
    arrow(ax, 50, 85, 50, 79)
    arrow(ax, 50, 68, 50, 62)
    arrow(ax, 32, 42, 24, 33)
    arrow(ax, 68, 42, 76, 33)
    ax.text(24, 38, '是', fontsize=9, color='#166534')
    ax.text(75, 38, '否', fontsize=9, color='#1e40af')
    arrow(ax, 24, 14, 35, 9)
    arrow(ax, 76, 14, 65, 9)
    fig.tight_layout()
    fig.savefig(out/'fig6_client_reconnect.png', bbox_inches='tight')
    plt.close(fig)
    return out


def build():
    source_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    doc = Document(SOURCE)
    replace(doc, '1. 分布式系统目标与透明性权衡：',
        '1. 分布式系统目标与透明性权衡：资源共享使不同终端能够共同操作地图、标记与操作历史；访问透明性隐藏服务接口背后的实现差异，位置与迁移透明性使房间换节点后仍可用稳定标识访问；复制透明性隐藏存储副本的位置。故障透明性有成本，故障切换或恢复时应向用户说明暂时不可写。透明性讨论的是隐藏分布细节，不意味着让玩家看到无权访问的数据；迷雾属于访问控制。开放性依靠标准接口，可扩展性通过房间间并行、异步通信和静态资源分发获得。[1]')
    replace(doc, '4. 持久化与基础设施层',
        '4. 持久化与基础设施层（Storage & Infrastructure Tier）：权威数据库保存房间状态、增量事件及幂等结果；路由缓存和在线状态属于可重建辅助信息。地图资源服务保管预置底图原文件，按会话角色及当前迷雾裁剪并返回授权瓦片，原文件不向玩家提供公开直链。公共 CDN 仅缓存公开、非敏感资源；含受限信息的瓦片须经鉴权且禁止共享缓存。')
    cell(table(doc, '1-1'), 2, 2, '主持人选定预置地图，成员按权限加载地图瓦片')
    cell(table(doc, '1-1'), 2, 3, '静态资源命名、受控瓦片访问、共享场景初始化')
    replace(doc, '1. 并发模型（Concurrency Models）：',
        '1. 并发模型：接入层以非阻塞 I/O 处理连接，空间计算和存储交给工作单元；同房间写入仍由单一队列串行完成，不同房间可并行。浏览器保留地图渲染等本地计算，属于职责偏向展示的客户端，并非完全无计算的终端。2. 虚拟化：虚拟机或容器提供部署隔离，不能代替状态一致性。3. 代码迁移：下载前端脚本可联系课件的弱代码移动；本项目不迁移运行栈和执行上下文，不需要强移动性，也不把房间数据迁移误称为进程代码迁移。服务器集群通过接入层路由分配房间；TCP handoff 涉及连接移交，当前通过重连重新定位而不依赖该机制。[1]')

    cell(table(doc, '4-1'), 2, 0, '预置地图授权瓦片查询')
    cell(table(doc, '4-1'), 2, 3, '鉴权后裁剪；受限瓦片不进公共缓存')
    cell(table(doc, '4-1'), 4, 2, '服务端逐成员单播，逻辑上组播')
    replace(doc, '结合图 3，移动标记的核心协作时序流转如下：',
        '结合图 3，一次移动经过提交意图、接入路由、权威校验、同库事务确认和授权传播五步。请求携带标记标识、目标位置、base_version 及 operationId；协调服务检查身份、权限和版本，将状态、房间提交序号 roomSeq 与幂等结果原子保存，再返回确认。后台分发器扫描持久化事件，避免提交后进程崩溃造成永久漏通知。分发时比较每名成员操作前后的授权视图，形成新增、更新或移除事件；完全不可见的变化不下发。客户端依据成员专属视图序号 viewSeq 应用更新，roomSeq 仅表示权威历史位置。')
    replace(doc, '3. 广播机制的边界与可靠性假定：',
        '3. 广播与补发边界：房间广播是服务端逐成员单播形成的逻辑组播，不假定公网提供原子广播或可靠组播。[1, 15] WebSocket 所用 TCP 在同一有效连接上提供有序字节传送；应用仍要处理断连、进程崩溃和队列溢出。roomSeq 是房间提交序号，权限过滤会使其合法跳号，客户端不据此判断丢失；每个授权视图流另有连续 viewSeq，只有实际下发的可见事件占号。客户端对同一流的 viewSeq 去重、按序确认；缺口、超时或心跳报告的流进度落后触发补发。补发记录缺失时转为快照，绝不发送隐藏事件填号。')
    before(doc, '5. 命名',
        '4. 通信模型的补充比较：RPC 通过客户端存根封装参数、服务端存根解包与调用、返回值反向传递，提供远程调用抽象，但远程调用仍有网络超时和部分故障；本设计管理接口采用 HTTP 请求—应答，不把 REST 与 RPC 概念等同。异步调用可避免等待通知阻塞，但不消除失败语义。消息队列可实现持久化中介；当前用数据库日志承接这一恢复职责，无需另设业务消息队列。课件面向流通信包含离散字节流及有时间约束的连续媒体；桌游使用离散状态消息，目标是应用延迟，不要求音视频抖动控制、流间同步或网络带宽 QoS。小房间适合服务端扇出，树形或流言传播更适合大规模群体，本阶段不采用。[1]')

    cell(table(doc, '5-1'), 3, 2, '幂等记录保留期内；绑定用户、房间及历史代次')
    cell(table(doc, '5-1'), 3, 3, '权威库唯一约束及结果记录查询')
    cell(table(doc, '5-1'), 5, 3, '鉴权地图资源服务；公共缓存只放非敏感内容')
    before(doc, '6. 同步',
        '4. 其他命名方式与解析：DNS 将服务域名解析为接入地址，接入层再以 roomId 查受控目录，形成层次化名称解析；转发指针易形成迁移链，家乡节点定位引入固定依赖，当前选择直接更新房间目录并重连定位。属性命名用于按地图标签、尺寸等条件检索，输出 mapId 后仍须鉴权，不以搜索结果代替授权；现阶段普通索引查询即可，无需部署 LDAP。DHT/Chord 适合开放对等查找，本系统受控节点规模下不采用。[1]')

    replace(doc, '结合图 4，对并发冲突的处理逻辑阐述如下：',
        '结合图 4，假设两名玩家均有权查看该标记：二者携带相同 base_version 请求移动，协调者依队列处理，首个版本匹配的请求完成事务并递增对象版本；随后请求因版本落后被拒绝，返回其有权查看的当前状态。冲突方撤销非权威预览。不同客户端之间由服务端决定仲裁次序；同一客户端的写意图则只允许一个未决请求，后续请求须待前一请求结果明确后发出。超时先用同一 operationId 查询或重试，未决之前不发依赖它的新写；重连仍保留这一顺序。这样房间提交序列既统一定序，也保留每个客户端的程序顺序。权限不同的成员只应用该权威历史的授权投影，不要求所有成员收到完全相同的数据。')
    replace(doc, '第一，明确区分 Lamport 逻辑时钟与服务端递增事件序号：',
        '第一，Lamport 时钟遵循因果先行 a→b 蕴含 C(a)<C(b)，反向不成立，单凭标量时钟不能判定两个事件是否并发；向量时钟能够识别因果与并发关系。[1, 3, 6] 本系统的 roomSeq 是集中权威提交序号，viewSeq 是成员授权流序号，都不冒称 Lamport 时钟。核心操作依赖服务端已确认的状态与对象版本，不承诺未提交客户端意图之间的完整因果追踪。')
    cell(table(doc, '6-1'), 2, 1, '满足因果先行则逻辑时钟递增；逆命题不成立')
    before(doc, '7. 一致性与复制',
        '第三，协调者选举与唯一写入权是两个问题：课件 Bully 算法由发现故障者联系编号更高的进程，最高可用编号者成为协调者；环选举沿逻辑环传播候选者，再公布结果。[1] 当前受控集群通过数据库中房间归属记录的原子条件更新选择接管者，候选者先获得新任期再恢复写入，而不让浏览器运行 Bully 或环选举。超时只触发接管尝试，不证明旧进程已停止；存储逐次校验 fencing 任期才是阻止旧写入的安全边界，拿到最高编号本身不能代替该隔离。')

    replace(doc, '（1）核心对局状态的确定性需求：',
        '（1）核心对局状态的确定性需求：标记、场景与掷骰结果形成单房间统一提交序列，保留各客户端已提交操作的程序顺序；各成员只观察与其权限相符的投影，不把网络到达时间视为全局因果证明。')
    replace(doc, '（3）非关键瞬态信息的最终收敛需求：',
        '（3）非关键瞬态信息的收敛需求：光标位置和输入提示允许短暂滞后，以最新采样覆盖旧状态；采样停止但成员仍在线时须重复发送最新状态直至收到确认，离线状态以会话失效清除，才能在网络恢复后收敛。')
    t = table(doc, '7-1')
    cell(t, 1, 4, '单房间串行提交且保留会话程序顺序；成员应用授权投影')
    cell(t, 3, 3, '同一历史代次内不退回已确认或已读水位之前')
    cell(t, 3, 4, '携带最大已读/已提交水位；不足时等待或拒绝；历史回退须显式重置')
    cell(t, 4, 4, '最新采样覆盖旧值；末次状态重发确认，离线用会话失效清除')
    replace(doc, '（1）顺序一致性（Sequential Consistency）：',
        '（1）顺序一致性要求操作结果可解释为同一串行序列，并保留每个进程的程序顺序。[1, 2] 系统以单房间串行事务和前述单未决写规则满足这一设计约束；读取来自该权威历史的已提交前缀，授权视图按序更新。前端推测状态不是已提交读取；网络延迟下不同成员可暂时处于不同前缀，系统不承诺所有端同时看到最新状态，也不把授权投影说成完整数据副本。')
    replace(doc, '（2）乐观交互预览（Optimistic UI Preview）与读自身写',
        '（2）乐观预览与客户端一致性：拖拽预览尚未提交，不能证明读自身写。确认后，客户端记录最大已提交及已读的权威水位，后续快照和读取不得低于它；存储尚未追齐时等待或返回暂不可读。[1, 12] 这一读自身写和单调读保证限于同一未回退的数据历史代次。若异步备库缺失已确认记录，单凭水位无法恢复数据，系统暂停该房间，等待补齐原日志；确实无法补齐时，须显式告知历史回退并经主持人确认后更换数据历史代次、重置全部视图。重置前不返回较旧结果冒充连续读取，旧代次的未决 operationId 不自动重放。')
    before(doc, '(2) 战争迷雾与成员权限的状态投影隔离',
        '（3）其余模型及适用边界：单调写要求同一客户端后续写入基于已经包含其先前写入的副本；写跟随读要求写入所用状态不落后于该客户端最近读取。[1] 本设计以同一权威路径、客户端水位和对象版本检查约束这两项要求；条件不满足时拒绝或同步，不能让旧视图直接覆盖。因果一致性强调有因果关系的写入按序观察，独立并发写可不同序；本项目选择权威全序，不额外提供全 P2P 因果追踪。连续一致性以 Conit 为测量单位，分别限制数值、陈旧度和顺序偏差：角色位置及掷骰结果不接受任意数值偏差；光标可容忍滞后，但当前未为其量化三个上界，故不宣称实现了有界连续一致性。')
    replace(doc, '结合图 5，权限隔离与状态投影的设计原则如下：',
        '结合图 5，服务端在实时事件、快照及地图瓦片上统一鉴权。预置底图原文件仅存服务端，玩家按地图标识、瓦片坐标及当前权限获取服务端裁剪后的图像；未探索像素及隐藏实体坐标、属性不进入玩家响应，公共 CDN 不保存受限瓦片。实体由不可见变为可见时下发当前授权状态，由可见变为不可见时发送只含已知标识的移除事件；双方均不可见的变化不占 viewSeq。成员权限或迷雾规则版本改变时终止旧视图流，发送重置通知，重新获取授权快照。已经披露的信息不能从玩家记忆或截图中撤回，撤权保证的是停止后续披露并清理当前视图，不作追溯保密承诺。')
    replace(doc, '1. 主备复制模式（Primary-Backup）：',
        '1. 主备复制由主节点受理写入并传播日志，适合本系统单权威路径。若采用异步复制，主库已确认但尚未复制的记录可能在灾难故障后丢失，恢复遵循前述历史代次规则；若要求切换后保留全部已确认写，则确认条件必须包含所需副本的持久化及安全切换前提，不能只改一句“零丢失”。2. Quorum 比较：等权 N 副本条件 R+W>N、W>N/2 保证读写及写写集合相交；加权模型以总票数 V 替代 N，要求读写票数阈值满足 R+W>V、W>V/2。[1, 4] 交集条件还需版本选择、并发协调及故障假设，不能单独推出任意网络分区下的强一致。当前应用层不自建投票协议，交给成熟存储设施。')
    before(doc, '8. 容错性与故障恢复',
        '3. 副本管理与传播：公开静态资源适合 CDN/客户端缓存，核心状态读取以权威库为准，不任意从落后备库读取。课件区分永久副本、服务器发起及客户端发起副本，更新可推送、拉取或混合；本系统在线推送、恢复时拉取，缓存只服务资源访问，不越过权限和水位。主动复制要求多个节点按同序确定性执行，随机掷骰还需先统一随机结果；本项目选择主节点生成并记录结果的主备方式。本地写/迁移主节点可优化访问局部性，但增加归属转移复杂度，小房间当前不采用。[1]')

    replace(doc, '针对上述网络不确定性，系统设计了',
        '针对上述不确定性，客户端为每个写意图生成 operationId，重试复用原标识及原载荷。服务端先验证当前身份与历史代次，再以用户、房间、代次和 operationId 的唯一约束，将状态变更及原结果原子保存；命中已提交记录只返回原有授权结果，相同标识而载荷不同则拒绝，仍在处理则稍后重试。记录须覆盖允许的重试期限；超过保留期或发生历史代次重置时，不把旧意图当成新意图自动执行，须重新同步并明确结果。这样约束的是保留期内、同一持久历史的重复副作用，而非无条件的恰好一次。')
    replace(doc, '结合图 6，系统对重连恢复的分支判定推导如下：',
        '结合图 6，重连先鉴权，客户端提交流标识、最后连续应用并确认的 viewSeq、权限版本和历史代次。同一授权流的补发记录完整时，服务端从确认点后补发可见事件；客户端拒绝旧流消息，对重复序号去重，对缺口暂停应用，不用 roomSeq 的合法跳号判漏收。授权视图流短期保留补发记录；服务实例重启而无法恢复原流、权限版本变化或记录超窗时，均改为快照重置。')
    before(doc, '表 8-1  ',
        '快照重置先在权威库固定一致的提交基线 B，读取对应授权状态，同时缓冲 B 之后的已提交变化。客户端整体替换旧视图以清除失效对象，获得新流标识并确认基线后，再连续应用缓冲的可见事件；若缓冲溢出、权限再次改变或代次不符，废弃本次快照并重试，不能直接切到实时监听留下中间缺口。后台扫描持久日志并重试派发，心跳核对可见流进度，覆盖提交后未广播及末条消息遗漏。本机制是权威状态截面恢复，不是 Chandy–Lamport 去中心化全局快照算法；文献仅用于对照通道状态与截面概念。[1, 11] 五秒恢复仍是未经测试的设计目标。')
    t = table(doc, '8-1')
    cell(t, 4, 0, '实时事件应用层遗漏或连接中断')
    cell(t, 4, 1, '网关队列溢出、进程异常或断连导致未确认消息缺失')
    cell(t, 4, 2, '检查同一授权流 viewSeq/心跳进度；补发不足则授权快照重置')
    cell(t, 3, 2, '原子获取新任期并核验 fencing；从持久化状态恢复；旧流不可续则快照重置')
    replace(doc, '当负责某房间的协作服务实例发生故障时，',
        '协调服务故障时，网关依据心跳超时仅判断“疑似失联”，由候选实例在权威存储上原子争取该房间的新任期，成功者加载已提交状态后接管。每次写入都在同一事务内检查有效任期，拒绝旧实例；续租与到期判定采用受控服务端时间或一致性存储，不信任客户端时钟。租约有时钟及暂停假设，安全边界最终是存储的原子任期检查；权威存储不可达则暂停写入。该方案是概要设计，尚未通过自动接管或网络分区测试，不把超时、最高编号或路由修改本身当作唯一写入证明。[1, 2]')
    replace(doc, '设计论证：2PC/3PC 适用于',
        '设计论证：2PC 面向多个独立事务参与者的原子提交，参与者在准备后若无法获知决定可能阻塞；3PC 增加预提交阶段，其非阻塞讨论依赖有界延迟等更强假设，网络分区或假设失效时不能据此保证安全且可用。[1, 10] 当前将房间状态、事件和幂等结果纳入同一权威库事务，数据库在正确持久化与隔离配置下承担 ACID，不涉及独立资源管理器的原子提交，故不自建 2PC/3PC。数据库复制及通知送达仍是另外的边界；单库事务不自动解决它们。')
    before(doc, '9. 总结',
        '4. 进程韧性、可靠组通信与恢复：课件 k 容错讨论在故障假设下允许 k 个节点失效所需的冗余；崩溃故障与拜占庭故障的副本数量条件不同，本项目未固定副本部署，不宣称达到某个 k 容错等级。房间成员加入、退出和权限变化可联系组视图及虚拟同步，但本系统采用权威服务生成不同授权流，不把该机制等同为所有组成员收到同一组消息的虚拟同步协议。持久事件日志、授权状态检查点与回放用于前向恢复，旧进程恢复后须重新核验任期；不回滚客户端画面就声称完成分布式故障恢复。[1, 15]')

    t = table(doc, '9-1')
    cell(t, 4, 1, 'RPC 存根、瞬态/持久消息、离散/连续流、组播')
    cell(t, 4, 2, 'HTTP 管理接口、双向消息、持久日志扫描和授权视图流')
    cell(t, 4, 3, '不把 REST 当作 RPC；无音视频，不引入媒体 QoS 和大规模组播')
    cell(t, 5, 1, '平面/结构化/属性命名、DHT、目录解析')
    cell(t, 5, 2, '稳定 ID、邀请码、受控路由、地图标签检索及鉴权瓦片')
    cell(t, 5, 3, '比较转发指针、家乡节点与 Chord；当前无需 LDAP 或覆盖网络')
    cell(t, 6, 1, '物理/逻辑时钟、向量时钟、互斥、Bully/环选举')
    cell(t, 6, 2, '单房间定序、会话程序顺序、版本检测与原子任期接管')
    cell(t, 6, 3, '房间序号不等于 Lamport 时钟；选举不等于 fencing')
    cell(t, 7, 1, '顺序/因果/最终/连续一致性、四类客户端模型、复制')
    cell(t, 7, 2, '授权投影、视图序号、水位、快照衔接与历史回退边界')
    cell(t, 7, 3, 'Conit 不承诺未量化界限；异步回退暂停并显式更换历史代次')
    cell(t, 8, 1, '故障模型、k 容错、可靠组通信、日志恢复、2PC/3PC')
    cell(t, 8, 2, '保留期内幂等、快照/增量分流、存储任期隔离')
    cell(t, 8, 3, '不宣称拜占庭或虚拟同步实现；单库不自建跨库提交')
    replace(doc, '纵观全篇报告，本设计围绕实际业务需求',
        '本设计将课程知识落实到并发移动、掷骰重试和迷雾隔离：C/S 统一校验，单房间定序与对象版本处理冲突，房间间并行支撑扩展。授权视图流区分隐藏事件与漏收，实体移除及快照维护正确视图，受控瓦片限制底图披露。对 CRDT、P2P、应用层 Quorum 和跨库提交，均依据当前需求说明不采用的理由。[1, 13]')
    replace(doc, '遵循严谨求实的科学态度，报告对本设计的局限性',
        '本报告是概要设计，未完成原型、压测或故障注入。每房间最多 8 人、约 100 个活跃房间、实时 P95≤300 ms、API P95≤500 ms、恢复≤5 秒均是设计目标。异步复制丢失已确认历史时须暂停并显式恢复，跨历史代次不承诺连续读自身写或去重；已披露信息不能撤回。后续须验证授权瓦片、视图流、fencing 和快照衔接，并评估单房间串行瓶颈，不能把当前推导当作已实现保证。')

    fill(find(doc, '[1] 宋杰.'), '[1] 宋杰（修改）. 分布式系统导论（Introduction to Distributed System）[Z]. 东北大学软件学院教学课件, 2014（据课件页脚修改年份著录）.', size=9, refs=True)
    fill(find(doc, '[9] Owlbear Rodeo.'), '[9] Owlbear Rodeo. Owlbear Rodeo [EB/OL]. https://www.owlbear.rodeo/（网址见课堂选题资料）.', size=9, refs=True)

    # Restore the exact supplied appendix table, then refill existing team data.
    old = doc.tables[-1]
    saved_values = [c.text.strip() for c in old.rows[2].cells]
    raw = Document(ROOT/'raw/课程设计模板.docx')
    restored_el = deepcopy(raw.tables[-1]._tbl)
    old._tbl.getparent().replace(old._tbl, restored_el)
    restored = Table(restored_el, doc)
    for i, c in enumerate(restored.rows[2].cells):
        fill(c.paragraphs[0], saved_values[i], size=10.5, refs=True)
        for p in c.paragraphs[1:]:
            p._p.getparent().remove(p._p)

    images = draw_diagrams()
    for fig_num, image_name in [(5, 'fig5_permission_projection.png'), (6, 'fig6_client_reconnect.png')]:
        cap = find(doc, f'图 {fig_num}  ')
        image_p = cap._p.getprevious()
        rel_id = image_p.xpath('.//a:blip')[0].get(qn('r:embed'))
        doc.part.related_parts[rel_id]._blob = (images/image_name).read_bytes()

    # Remove Markdown code delimiters in prose and keep body line spacing explicit.
    for p in doc.paragraphs:
        if '`' in p.text:
            fill(p, p.text.replace('`', ''))
        if p.style.name == 'Normal' and len(p.text) > 80:
            spacing = p._p.get_or_add_pPr().find(qn('w:spacing'))
            if spacing is not None:
                spacing.set(qn('w:beforeLines'), '15')
                spacing.set(qn('w:afterLines'), '15')

    settings = doc.settings.element
    uf = settings.find(qn('w:updateFields'))
    if uf is None:
        uf = OxmlElement('w:updateFields')
        settings.append(uf)
    uf.set(qn('w:val'), 'true')
    # All ten content tables fit a page. Keep their rows together without fixed page numbers.
    for number in ['1-1', '1-2', '2-1', '3-1', '4-1', '5-1', '6-1', '7-1', '8-1', '9-1']:
        t = table(doc, number)
        for i, row in enumerate(t.rows):
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.keep_with_next = i < len(t.rows)-1
    find(doc, '参考文献').paragraph_format.page_break_before = True
    for p in doc.paragraphs:
        if re.match(r'^\[\d+\]', p.text):
            p.paragraph_format.keep_together = True
    doc.core_properties.title = '雾桌（MistTable）多人在线虚拟桌游平台的技术分析和概要设计'
    doc.core_properties.revision = 3
    doc.save(DOCX)
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == source_hash
    (WORK/'build-summary.json').write_text(json.dumps({
        'source': str(SOURCE), 'source_sha256': source_hash,
        'docx': str(DOCX), 'paragraphs': len(doc.paragraphs),
        'tables': len(doc.tables), 'images': len(doc.inline_shapes),
        'appendix_restored_from_template': True,
    }, ensure_ascii=False, indent=2), encoding='utf-8')
    print('Built:', DOCX)


if __name__ == '__main__':
    build()
