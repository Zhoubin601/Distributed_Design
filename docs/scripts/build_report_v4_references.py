"""Apply the reviewed reference corrections to v3 without changing its structure."""
from pathlib import Path
from copy import deepcopy
from hashlib import sha256
import json
import re
import sys

from docx import Document
from docx.shared import Pt
from scratch_report_builder_core import parse_body_text_segments

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT/'output/v3/20261024-misttable-distributed-design-v3.docx'
WORK = ROOT/'docs/20261024-reference-audit-v1'
TARGET = WORK/'20261024-misttable-distributed-design-v4.docx'

REPLACEMENTS = [
    ('系统采用分层客户端—服务器（Layered Client-Server）架构与事件通知相结合的风格',
     '系统采用分层客户端—服务器（Layered Client-Server）架构。[1, 7] 在此基础上，本设计用事件通知传播房间状态。系统整体划分为客户端层、接入层、应用与房间协作服务层、持久化与基础设施层四个逻辑层次。系统总体逻辑架构如图 1 所示。'),
    ('系统采用‘请求—应答（Request-Reply）’与‘双向实时通信’相结合的通信架构',
     '系统将常规管理的请求—应答与实时协作的双向消息通道分开设计。[1] WebSocket 提供基于 TCP 的双向通信，可用于后者。[8] 不同业务场景下的通信方式对比如表 4-1 所示。'),
    ('系统采用全局唯一标识符（如 UUIDv4）作为业务实体的稳定内部名称',
     '命名设计遵循稳定标识与网络位置分离的原则。[1, 2] 本系统选择 UUIDv4 作为业务实体的内部标识，并由权威库唯一约束拒绝重复值。核心实体包括用户标识（userId）、房间内部标识（roomId）、标记标识（tokenId）与地图资源标识（mapId）。UUIDv4 用于内部引用与寻址，其随机性不能替代授权；获取某个实体标识不等于获得操作权限，服务端仍须核验会话身份、成员角色和对象权限。'),
    ('深入理论剖析与概念澄清：课件详细证明了 Lamport 逻辑时钟',
     '理论概念对应：Lamport 逻辑时钟 [3] 说明因果先行与时间戳的关系，向量时钟 [6] 用于识别因果与并发，Ricart-Agrawala 算法 [5] 解决消息传递网络中的互斥。以下将理论模型与本系统工程机制区分说明：'),
    ('协调服务故障时，网关依据心跳超时仅判断',
     '依据课程的故障检测与恢复模型 [1, 2]，本设计在协调服务故障时仅将心跳超时视为“疑似失联”。候选实例在权威存储上原子争取该房间的新任期，成功者加载已提交状态后接管。每次写入都在同一事务内检查有效任期，拒绝旧实例；续租与到期判定采用受控服务端时间或一致性存储，不信任客户端时钟。租约有时钟及暂停假设，安全边界最终是存储的原子任期检查；权威存储不可达则暂停写入。该方案尚未通过自动接管或网络分区测试，不把超时、最高编号或路由修改本身当作唯一写入证明。'),
    ('3. 分布式两阶段提交（2PC）与三阶段提交（3PC）适用性论证：课件 Chapter 8',
     '3. 分布式两阶段提交（2PC）与三阶段提交（3PC）适用性论证：课件 Chapter 8 及教材第 8.5 节介绍了 2PC 与 3PC 在保障跨节点事务原子提交方面的机制与限制。[1, 2]'),
    ('设计论证：2PC 面向多个独立事务参与者的原子提交',
     '设计论证：2PC 面向多个独立事务参与者的原子提交，参与者在准备后若无法获知决定可能阻塞。[1, 10] 3PC 增加预提交阶段，其非阻塞讨论依赖有界延迟等更强假设，网络分区或假设失效时不能据此保证安全且可用。[1, 2] 当前将房间状态、事件和幂等结果纳入同一权威库事务，数据库在正确持久化与隔离配置下承担 ACID，不涉及独立资源管理器的原子提交，故不自建 2PC/3PC。数据库复制及通知送达仍是另外的边界；单库事务不自动解决它们。'),
    ('4. 进程韧性、可靠组通信与恢复：课件 k 容错讨论',
     '4. 进程韧性、可靠组通信与恢复：课件 k 容错讨论在故障假设下允许 k 个节点失效所需的冗余；崩溃故障与拜占庭故障的副本数量条件不同，本项目未固定副本部署，不宣称达到某个 k 容错等级。[1] 文献 [15] 讨论故障条件下可靠组通信及事件交付顺序；本系统采用不同授权视图流，不将其等同于虚拟同步协议。恢复时先回到已提交的检查点，再回放持久日志，属于检查点结合日志的后向恢复。[1, 2] 旧进程恢复后须重新核验任期，不以客户端画面回滚代替分布式故障恢复。'),
    ('[5] RICART G, AGRAWALA A K.',
     '[5] RICART G, AGRAWALA A K. An optimal algorithm for mutual exclusion in computer networks [J]. Communications of the ACM, 1981, 24(1): 9-17.'),
    ('[6] MATTERN F.',
     '[6] MATTERN F. Virtual time and global states of distributed systems [C]//Proceedings of the Workshop on Parallel and Distributed Algorithms. North-Holland, 1989: 215-226.'),
    ('[9] Owlbear Rodeo.',
     '[9] Owlbear Rodeo. Owlbear Rodeo [EB/OL]. [2026-10-24]. https://www.owlbear.rodeo/.'),
    ('[10] GRAY J.',
     '[10] GRAY J. Notes on data base operating systems [M]//Operating Systems: An Advanced Course. Lecture Notes in Computer Science, vol. 60. Berlin: Springer, 1978: 393-481.'),
    ('[13] SHAPIRO M,',
     '[13] SHAPIRO M, PREGUIÇA N, BAQUERO C, et al. Conflict-free replicated data types [C]//Stabilization, Safety, and Security of Distributed Systems. SSS 2011. Lecture Notes in Computer Science, vol. 6976. Berlin: Springer, 2011: 386-400.'),
]


def fill(p, text, bibliography=False):
    sample = next((r for r in p.runs if r.text.strip()), None)
    props = deepcopy(sample._r.rPr) if sample is not None and sample._r.rPr is not None else None
    p.clear()
    for value, sup in ([(text, False)] if bibliography else parse_body_text_segments(text)):
        r = p.add_run(value)
        if props is not None:
            r._r.insert(0, deepcopy(props))
        r.font.superscript = sup
        if bibliography:
            r.font.size = Pt(9)


def build():
    WORK.mkdir(exist_ok=True)
    if TARGET.exists() and '--rebuild' not in sys.argv:
        raise FileExistsError(f'Preserve existing version: {TARGET}')
    source_hash = sha256(SOURCE.read_bytes()).hexdigest()
    doc = Document(SOURCE)
    contexts = []
    for i, p in enumerate(doc.paragraphs):
        if re.search(r'\[\d+(?:,\s*\d+)*\]',p.text):
            contexts.append({'paragraph': i+1, 'text': p.text})
    for ti,t in enumerate(doc.tables):
        for ri,row in enumerate(t.rows):
            for ci,c in enumerate(row.cells):
                if re.search(r'\[\d+(?:,\s*\d+)*\]',c.text):
                    contexts.append({'table':ti+1,'row':ri+1,'column':ci+1,'text':c.text})
    (WORK/'v3-citation-contexts.json').write_text(json.dumps(contexts,ensure_ascii=False,indent=2),encoding='utf-8')
    changes = []
    for prefix,new in REPLACEMENTS:
        matches = [p for p in doc.paragraphs if p.text.startswith(prefix)]
        if len(matches)!=1:
            raise ValueError(f'{prefix}: {len(matches)} matches')
        p=matches[0]
        old=p.text
        fill(p,new,bibliography=prefix.startswith('['))
        changes.append({'before':old,'after':new})
    # Do not imply that every P2P design must use consensus or a CRDT.
    table_changes=0
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    if p.text=='节点间对等同步，缺少中心权威，需依赖共识算法或 CRDT [13] 消解冲突':
                        old=p.text
                        new='需另行设计定序与合并；CRDT [13] 仅保证满足相应条件的数据类型收敛'
                        fill(p,new)
                        changes.append({'before':old,'after':new})
                        table_changes+=1
    assert table_changes==1
    doc.core_properties.revision=4
    doc.save(TARGET)
    assert sha256(SOURCE.read_bytes()).hexdigest()==source_hash
    (WORK/'change-manifest.json').write_text(json.dumps({
        'date':'2026-10-24','source':str(SOURCE),'source_sha256':source_hash,
        'target':str(TARGET),'changes':changes,
        'paragraph_count':len(doc.paragraphs),'table_count':len(doc.tables),
        'inline_image_count':len(doc.inline_shapes)},ensure_ascii=False,indent=2),encoding='utf-8')
    print(f'Created {TARGET}; {len(changes)} targeted changes.')


if __name__=='__main__':
    build()
