"""Read-only report audit; evidence stays under docs, personal data is not logged."""
from pathlib import Path
from hashlib import sha256
from collections import Counter
import json
import re
import subprocess
from zipfile import ZipFile

import fitz
from PIL import Image, ImageDraw
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
EVIDENCE = ROOT / 'docs/20261024-v3-recheck-v1'
EVIDENCE.mkdir(exist_ok=True)
REPORT = ROOT / 'output/v3/20261024-misttable-distributed-design-v3.docx'
PDF = REPORT.with_suffix('.pdf')

def digest(path):
    return sha256(path.read_bytes()).hexdigest()

before = {str(p.relative_to(ROOT)): digest(p)
          for base in ('raw', 'output') for p in (ROOT/base).rglob('*') if p.is_file()}
summary = {'date': '2026-10-24', 'head': subprocess.check_output(
    ['git', 'rev-parse', 'HEAD'], cwd=ROOT, text=True).strip(),
    'origin_main': subprocess.check_output(
        ['git', 'rev-parse', 'origin/main'], cwd=ROOT, text=True).strip()}

# Read the exact remote report without restoring old paths in the working tree.
remote_bytes = subprocess.check_output(
    ['git', 'show', 'origin/main:output/Distributed_Design_Report_Draft_v2.docx'], cwd=ROOT)
remote_path = EVIDENCE/'remote-baseline-v2.docx'
remote_path.write_bytes(remote_bytes)
remote = Document(remote_path)
local = Document(REPORT)
template = Document(next((ROOT/'raw').glob('*.docx')))
previous = Document(ROOT/'output/v2/Distributed_Design_Report_Draft_v2.docx')

def all_text(doc):
    return '\n'.join([p.text for p in doc.paragraphs] +
                     [c.text for t in doc.tables for row in t.rows for c in row.cells])

def grid_signature(table):
    return [[(c.grid_span, c.vMerge) for c in row._tr.tc_lst] for row in table.rows]

def unique_row_values(table):
    return [[''.join(c.xpath('.//w:t/text()')) for c in row._tr.tc_lst] for row in table.rows]

summary['reports'] = {
    'local_docx_sha256': digest(REPORT), 'local_pdf_sha256': digest(PDF),
    'remote_docx_sha256': sha256(remote_bytes).hexdigest(),
    'local_equals_remote': REPORT.read_bytes() == remote_bytes,
    'remote_placeholder_occurrences': len(re.findall('待填写|待确认|TODO', all_text(remote))),
    'local_placeholder_occurrences': len(re.findall('待填写|待确认|TODO', all_text(local))),
    'local_cover_same_as_manual_v2': unique_row_values(local.tables[0]) == unique_row_values(previous.tables[0]),
    'local_appendix_values_same_as_manual_v2': unique_row_values(local.tables[-1]) == unique_row_values(previous.tables[-1]),
    'appendix_grid_and_merges_match_template': grid_signature(local.tables[-1]) == grid_signature(template.tables[-1]),
    'appendix_no_added_paragraphs': [len(c.p_lst) for row in local.tables[-1].rows for c in row._tr.tc_lst]
        == [len(c.p_lst) for row in template.tables[-1].rows for c in row._tr.tc_lst],
}
summary['body_paragraph_properties'] = dict(Counter(
    str(p._p.pPr.xml) if p._p.pPr is not None else 'inherit'
    for p in local.paragraphs[67:260] if p.style.name.lower() == 'normal'))
with ZipFile(REPORT) as z:
    summary['document_defaults_xml'] = z.read('word/styles.xml').decode('utf-8').split('</w:docDefaults>')[0].split('<w:docDefaults>')[-1]

pdf = fitz.open(PDF)
summary['pdf'] = {'pages': len(pdf), 'page_sizes': [[p.rect.width, p.rect.height] for p in pdf],
                  'text_chars_per_page': [len(p.get_text().strip()) for p in pdf]}
summary['pdf']['font_summary'] = dict(Counter(
    (span['font'] + ':' + str(span['size'])) for i in range(3,25)
    for block in pdf[i].get_text('dict')['blocks'] if block['type'] == 0
    for line in block['lines'] for span in line['spans']))
for group_start in range(0, len(pdf), 9):
    sheet = Image.new('RGB', (1500, 2220), '#d8d8d8')
    draw = ImageDraw.Draw(sheet)
    for k, page_index in enumerate(range(group_start, min(group_start+9, len(pdf)))):
        pix = pdf[page_index].get_pixmap(matrix=fitz.Matrix(0.8,0.8), alpha=False)
        thumb = Image.frombytes('RGB', (pix.width,pix.height), pix.samples)
        thumb.thumbnail((480,700))
        x, y = (k%3)*500+10, (k//3)*740+30
        sheet.paste(thumb,(x,y))
        draw.text((x,y-20), f'Page {page_index+1}', fill='black')
    sheet.save(EVIDENCE/f'contact-{group_start//9+1}.png')
for page_number in (1,12,16,19,21,24,25,26):
    pdf[page_number-1].get_pixmap(matrix=fitz.Matrix(1.5,1.5), alpha=False).save(
        EVIDENCE/f'page-{page_number:02}.png')

course = fitz.open(ROOT/'raw/Introduction to Distributed System-by JieSong NEU.pdf')
summary['course_pages'] = len(course)
selected = (1,155,261,319,321,346,347,348,349,350,370)
(EVIDENCE/'local-course-excerpts.txt').write_text('\n\n'.join(
    f'Physical page {p}\n'+course[p-1].get_text() for p in selected), encoding='utf-8')
after = {str(p.relative_to(ROOT)): digest(p)
         for base in ('raw', 'output') for p in (ROOT/base).rglob('*') if p.is_file()}
summary['raw_and_output_unchanged_during_audit'] = before == after
summary['input_hashes'] = before
(EVIDENCE/'independent-checks.json').write_text(
    json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({k: summary[k] for k in (
    'head','origin_main','reports','course_pages','raw_and_output_unchanged_during_audit')}, ensure_ascii=False,indent=2))
