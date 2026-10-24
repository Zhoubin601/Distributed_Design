"""Read-only audit of the delivered v4; evidence stays under docs/."""
from pathlib import Path
from hashlib import sha256
from zipfile import ZipFile
from collections import Counter
import json
import re
import fitz
from lxml import etree
from docx import Document
from PIL import Image, ImageDraw

ROOT = Path(__file__).resolve().parents[2]
WORK = ROOT / 'docs/20261024-v4-recheck-v1'
WORK.mkdir(exist_ok=True)
REPORT = ROOT / 'output/v4/20261024-misttable-distributed-design-v4.docx'
PDF = REPORT.with_suffix('.pdf')
def hashes():
    return {str(p.relative_to(ROOT)): sha256(p.read_bytes()).hexdigest()
            for base in ('raw', 'output') for p in (ROOT / base).rglob('*') if p.is_file()}
baseline = hashes()
(WORK / 'input-hashes.json').write_text(json.dumps(baseline, ensure_ascii=False, indent=2), encoding='utf-8')
doc = Document(REPORT)
template = Document(ROOT / 'raw/课程设计模板.docx')
ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
with ZipFile(ROOT / 'raw/测验方法.pptx') as z:
    slides = sorted((n for n in z.namelist() if re.match(r'ppt/slides/slide\d+\.xml$', n)),
                    key=lambda n: int(re.search(r'slide(\d+)', n).group(1)))
    (WORK / 'requirements-slides.txt').write_text('\n\n'.join(
        n + '\n' + '\n'.join(etree.fromstring(z.read(n)).xpath('//a:t/text()', namespaces=ns))
        for n in slides), encoding='utf-8')
def ordered_text(d):
    return '\n'.join(''.join(t.text or '' for t in el.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t')) for el in d.element.body)
(WORK / 'template.txt').write_text(ordered_text(template), encoding='utf-8')
(WORK / 'report.txt').write_text(ordered_text(doc), encoding='utf-8')
pdf = fitz.open(PDF)
(WORK / 'pdf-text.txt').write_text('\n\n'.join(
    f'=== PDF page {i+1} ===\n' + p.get_text() for i,p in enumerate(pdf)), encoding='utf-8')
def structure(table):
    return [[{'gridSpan': c.grid_span,
              'vMerge': c.vMerge,
              'paragraphs': len(c.p_lst)} for c in row._tr.tc_lst] for row in table.rows]
def values(table):
    return [[''.join(c.xpath('.//w:t/text()')) for c in row._tr.tc_lst] for row in table.rows]
v3 = Document(ROOT / 'output/v3/20261024-misttable-distributed-design-v3.docx')
manifest = json.loads((ROOT / 'docs/20261024-reference-audit-v1/change-manifest.json').read_text('utf-8'))
body_paragraphs = list(doc.paragraphs) + [p for t in doc.tables for row in t.rows for cell in row.cells for p in cell.paragraphs]
references = [p for p in doc.paragraphs if re.match(r'^\[\d+\]',p.text)]
checks = {
    'appendix_structure_matches_template': structure(doc.tables[-1]) == structure(template.tables[-1]),
    'appendix_grid_columns_match_template': len(doc.tables[-1]._tbl.tblGrid.gridCol_lst) == len(template.tables[-1]._tbl.tblGrid.gridCol_lst),
    'team_cover_unchanged_from_v3': values(doc.tables[0]) == values(v3.tables[0]),
    'team_appendix_unchanged_from_v3': values(doc.tables[-1]) == values(v3.tables[-1]),
    'no_placeholder_in_any_body_or_table': not re.search(r'待确认|待填写|TODO|lorem|xxxx', ordered_text(doc), re.I),
    'no_tracked_insertions_or_deletions': not doc.element.body.xpath('.//w:ins | .//w:del'),
    'pdf_pages_have_text': all(len(p.get_text().strip()) > 20 for p in pdf),
    'all_14_reference_revisions_in_delivered_v4': len(manifest['changes']) == 14 and all(any(p.text == c['after'] for p in body_paragraphs) for c in manifest['changes']),
    'reference_9_has_access_date': any(p.text.startswith('[9] ') and '[2026-10-24]' in p.text for p in references),
    'reference_count_15': len(references) == 15,
    'reference_font_size_9pt': all(r.font.size is not None and r.font.size.pt == 9 for p in references for r in p.runs if r.text.strip()),
}
with ZipFile(REPORT) as z:
    checks['native_page_field'] = any(b'PAGE' in z.read(n) for n in z.namelist() if re.match('word/footer.*xml$', n))
    checks['no_comments'] = not any(re.match('word/comments.*xml$', n) for n in z.namelist())
summary = {'checks': checks, 'failed': [k for k,v in checks.items() if not v],
           'pages': len(pdf), 'figures': len(doc.inline_shapes), 'tables': len(doc.tables),
           'headings': [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith('Heading')],
           'paragraph_count': len(doc.paragraphs),
           'normal_style_xml': doc.styles['Normal'].element.xml,
           'normal_paragraph_spacing_xml_counts': dict(Counter(str(p._p.pPr.xml) if p._p.pPr is not None else 'inherit' for p in doc.paragraphs if p.style.name == 'Normal')),
           'font_sizes': dict(Counter(s['font'] + ':' + str(round(s['size'], 1))
                for p in pdf for b in p.get_text('dict')['blocks'] if b['type'] == 0
                for line in b['lines'] for s in line['spans'])),
           'pages_text_chars': [len(p.get_text().strip()) for p in pdf]}
for start in range(0,len(pdf),9):
    sheet = Image.new('RGB',(1500,2220),'#d8d8d8')
    draw = ImageDraw.Draw(sheet)
    for k,i in enumerate(range(start,min(start+9,len(pdf)))):
        pix = pdf[i].get_pixmap(matrix=fitz.Matrix(.8,.8),alpha=False)
        thumb = Image.frombytes('RGB',(pix.width,pix.height),pix.samples)
        thumb.thumbnail((480,700))
        x,y = (k%3)*500+10,(k//3)*740+30
        sheet.paste(thumb,(x,y))
        draw.text((x,y-20),f'Page {i+1}',fill='black')
    sheet.save(WORK / f'contact-{start//9+1}.png')
for i in (1,12,17,19,21,22,23,25,26):
    pdf[i-1].get_pixmap(matrix=fitz.Matrix(1.5,1.5),alpha=False).save(WORK / f'page-{i:02}.png')
course = fitz.open(ROOT / 'raw/Introduction to Distributed System-by JieSong NEU.pdf')
selected = (1,155,261,321,346,347,348,349,350,370,498,499,500,501)
(WORK / 'local-course-excerpts.txt').write_text('\n\n'.join(f'Physical page {i}\n'+course[i-1].get_text() for i in selected),encoding='utf-8')
summary['course_page_count'] = len(course)
for extracted in WORK.glob('*.txt'):
    extracted.write_text('\n'.join(line.rstrip() for line in extracted.read_text('utf-8').splitlines()).rstrip()+'\n', encoding='utf-8')
summary['checks']['raw_and_output_unchanged'] = baseline == hashes()
(WORK / 'independent-checks.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k: summary[k] for k in ('checks','failed','pages','figures','tables','course_page_count')},ensure_ascii=False,indent=2))
if summary['failed'] or not summary['checks']['raw_and_output_unchanged']:
    raise SystemExit(1)
