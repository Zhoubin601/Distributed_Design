# -*- coding: utf-8 -*-
"""
主装配脚本：执行全量课程设计报告生成
"""
import os
import shutil
import docx

from scratch_report_builder_core import (
    set_anchor, set_run_font, set_p_spacing, add_heading_1, add_heading_2, add_heading_3,
    add_body_p, add_figure, add_table_with_caption
)
from scratch_report_builder_sections import (
    update_cover_page, update_toc
)
from scratch_build_ch1 import build_chapter_1
from scratch_build_ch2_ch3 import build_chapter_2, build_chapter_3
from scratch_build_ch4_ch5 import build_chapter_4, build_chapter_5
from scratch_build_ch6_ch7 import build_chapter_6, build_chapter_7
from scratch_build_ch8_ch9_refs import (
    build_chapter_8, build_chapter_9, build_references, update_appendix_table
)

TEMPLATE_PATH = 'raw/课程设计模板.docx'
OUTPUT_DIR = 'output'
OUTPUT_FILE_1 = os.path.join(OUTPUT_DIR, 'Distributed_Design_Report_Draft_v1.docx')
OUTPUT_FILE_2 = os.path.join(OUTPUT_DIR, '20261023-Distributed_Design_Report_Draft-v1.docx')

def main():
    print(">>> 正在加载原始课程设计模板...")
    doc = docx.Document(TEMPLATE_PATH)

    # 1. 更新封面信息与表 0
    print(">>> 正在更新封面标题与团队信息表...")
    update_cover_page(doc)

    # 2. 更新目录条目
    print(">>> 正在更新目录结构...")
    update_toc(doc)

    # 3. 定位附录起始锚点并清理旧模板占位段落
    print(">>> 正在定位附录起始锚点并清理旧模板占位段落...")
    body = doc.element.body
    elements = list(body)

    appendix_start_idx = 138
    for i, el in enumerate(elements):
        if el.tag.endswith('p'):
            p = docx.text.paragraph.Paragraph(el, doc)
            if '附录 (得分栏勿填)' in p.text:
                appendix_start_idx = i
                break

    anchor = elements[appendix_start_idx]
    old_body_nodes = elements[15:appendix_start_idx]

    # 清理旧的正文占位节点
    for node in old_body_nodes:
        body.remove(node)

    # 设置全局插入锚点
    set_anchor(anchor)

    # 4. 依次构建第 1 至第 9 章及参考文献
    print(">>> 正在构建第 1 章：系统简介...")
    build_chapter_1(doc)

    print(">>> 正在构建第 2 章：体系结构...")
    build_chapter_2(doc)

    print(">>> 正在构建第 3 章：进程...")
    build_chapter_3(doc)

    print(">>> 正在构建第 4 章：通信...")
    build_chapter_4(doc)

    print(">>> 正在构建第 5 章：命名...")
    build_chapter_5(doc)

    print(">>> 正在构建第 6 章：同步...")
    build_chapter_6(doc)

    print(">>> 正在构建第 7 章：一致性与复制...")
    build_chapter_7(doc)

    print(">>> 正在构建第 8 章：容错性与故障恢复...")
    build_chapter_8(doc)

    print(">>> 正在构建第 9 章：总结...")
    build_chapter_9(doc)

    print(">>> 正在构建参考文献...")
    build_references(doc)

    # 解除锚点锁定
    set_anchor(None)

    # 5. 更新附录表格信息
    print(">>> 正在更新附录考核表数据...")
    update_appendix_table(doc)

    # 6. 保存文件
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    doc.save(OUTPUT_FILE_1)
    shutil.copyfile(OUTPUT_FILE_1, OUTPUT_FILE_2)
    print(f"\n>>> 成功生成第一版 Word 草稿：\n    1. {OUTPUT_FILE_1}\n    2. {OUTPUT_FILE_2}")

    # 7. 快速自检
    doc_check = docx.Document(OUTPUT_FILE_1)
    print(f"\n>>> 文档自检结果：")
    print(f"    - 总段落数：{len(doc_check.paragraphs)}")
    print(f"    - 总表格数：{len(doc_check.tables)}")
    print(f"    - 分节数：{len(doc_check.sections)}")
    for i, s in enumerate(doc_check.sections):
        print(f"      Section {i}: 页面方向={s.orientation}, 宽={s.page_width.pt:.1f}pt, 高={s.page_height.pt:.1f}pt")

if __name__ == '__main__':
    main()
