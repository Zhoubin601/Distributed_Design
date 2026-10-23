# -*- coding: utf-8 -*-
"""
核心构建工具：段落、标题、三线表、图表插入辅助
支持在指定锚点（anchor）前插入，确保多节（横向/纵向）排版绝对稳定
"""

import os
import re
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

IMAGE_DIR = 'assets/diagrams/images'

# 当前插入锚点全局变量
CURRENT_ANCHOR = None

def set_anchor(anchor):
    global CURRENT_ANCHOR
    CURRENT_ANCHOR = anchor

def insert_p(doc, style=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    if CURRENT_ANCHOR is not None:
        CURRENT_ANCHOR.addprevious(p._p)
    return p

def insert_table(doc, rows, cols):
    tbl = doc.add_table(rows=rows, cols=cols)
    if CURRENT_ANCHOR is not None:
        CURRENT_ANCHOR.addprevious(tbl._tbl)
    return tbl

def set_run_font(run, east_asia='SimSun', ascii_font='Times New Roman', size_pt=10.5, bold=False, italic=False, color_rgb=None, superscript=False):
    run.font.name = ascii_font
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb
    if superscript:
        run.font.superscript = True
    rPr = run._r.get_or_add_rPr()
    if superscript and rPr.find(qn('w:noProof')) is None:
        rPr.append(OxmlElement('w:noProof'))
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:ascii'), ascii_font)
    rFonts.set(qn('w:hAnsi'), ascii_font)
    rFonts.set(qn('w:eastAsia'), east_asia)
    rFonts.set(qn('w:cs'), ascii_font)

def set_p_spacing(p, before=46, after=46, line=240, line_rule="auto"):
    pPr = p._p.get_or_add_pPr()
    spacing = pPr.find(qn('w:spacing'))
    if spacing is None:
        spacing = OxmlElement('w:spacing')
        pPr.append(spacing)
    spacing.set(qn('w:before'), str(before))
    spacing.set(qn('w:after'), str(after))
    spacing.set(qn('w:line'), str(line))
    spacing.set(qn('w:lineRule'), line_rule)

def add_heading_1(doc, text):
    p = insert_p(doc, style='Heading 1')
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_p_spacing(p, before=180, after=100)
    p.text = ''
    r = p.add_run(text)
    set_run_font(r, east_asia='SimHei', size_pt=16, bold=True)
    return p

def add_heading_2(doc, text):
    p = insert_p(doc, style='Heading 2')
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_p_spacing(p, before=140, after=60)
    p.text = ''
    r = p.add_run(text)
    set_run_font(r, east_asia='SimHei', size_pt=14, bold=True)
    return p

def add_heading_3(doc, text):
    p = insert_p(doc, style='Heading 3')
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_p_spacing(p, before=80, after=40)
    p.text = ''
    r = p.add_run(text)
    set_run_font(r, east_asia='SimHei', size_pt=12, bold=True)
    return p

CITATION_REGEX = re.compile(r'(\[\d+(?:,\s*\d+)*\])')

def parse_body_text_segments(text, citations=None):
    """
    解析正文中的文本与上标引用片段：
    - 自动识别正文中形如 [1]、[1, 7]、[9]、[1, 2] 的文献编号，并转换为上标引用；
    - 若传入兼容参数 citations=[(marker, ref_str), ...]，先安全对齐锚点，避免产生 [9][9] 重复。
    返回: list of (seg_text, is_superscript)
    """
    normalized_text = text
    if citations:
        for marker, ref_str in citations:
            if marker in normalized_text:
                # 检查 marker 是否已包含引用方括号
                if re.search(r'\[\d+(?:,\s*\d+)*\]', marker):
                    # marker 中本身已书写引用，无需重复追加
                    pass
                else:
                    # 检查 normalized_text 中 marker 之后是否紧跟 ref_str
                    idx = normalized_text.find(marker)
                    after = normalized_text[idx + len(marker):]
                    if not after.startswith(ref_str):
                        normalized_text = normalized_text[:idx + len(marker)] + ref_str + after

    tokens = CITATION_REGEX.split(normalized_text)
    segments = []
    for token in tokens:
        if not token:
            continue
        if CITATION_REGEX.fullmatch(token):
            segments.append((token, True))
        else:
            segments.append((token, False))
    return segments

def add_body_p(doc, text, citations=None):
    """
    添加标准正文段落：五号宋体，段前段后0.15行，单倍行距。首行缩进2字符。
    自动将 [1]、[1, 7] 等文献标记格式化为统一的上标引用。
    """
    p = insert_p(doc, style='Normal')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_p_spacing(p, before=46, after=46, line=240, line_rule="auto")
    pPr = p._p.get_or_add_pPr()
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind')
        pPr.append(ind)
    ind.set(qn('w:firstLineChars'), '200') # 首行缩进2字符

    segments = parse_body_text_segments(text, citations=citations)
    for seg_text, is_sup in segments:
        r = p.add_run(seg_text)
        set_run_font(r, east_asia='SimSun', size_pt=10.5, superscript=is_sup)
    return p

def add_figure(doc, img_name, fig_num, caption):
    """
    插入图片与图注：图片居中，图注在图下方居中，仿宋五号加粗。
    为图片段落设置 keep_with_next = True 避免图片与图注跨页分离。
    """
    img_path = os.path.join(IMAGE_DIR, img_name)
    if os.path.exists(img_path):
        p_img = insert_p(doc)
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_p_spacing(p_img, before=80, after=40)
        p_img.paragraph_format.keep_with_next = True
        r_img = p_img.add_run()
        r_img.add_picture(img_path, width=Inches(6.0))
        
        p_cap = insert_p(doc)
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_p_spacing(p_cap, before=20, after=80)
        r_cap = p_cap.add_run(f"图 {fig_num}  {caption}")
        set_run_font(r_cap, east_asia='FangSong', size_pt=10.5, bold=True)
    else:
        print(f"Warning: Image {img_path} not found!")

def add_table_with_caption(doc, tab_num, caption, headers, rows_data, col_widths=None):
    """
    添加三线表与表注：表注在表格上方居中，仿宋五号加粗；三线表顶底粗线、表头下细线。
    为表题设置 keep_with_next = True，确保表题绝不孤立滞留在前页底部，与后续表格始终保持在一起。
    """
    p_cap = insert_p(doc)
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_p_spacing(p_cap, before=80, after=30)
    p_cap.paragraph_format.keep_with_next = True
    r_cap = p_cap.add_run(f"表 {tab_num}  {caption}")
    set_run_font(r_cap, east_asia='FangSong', size_pt=10.5, bold=True)

    table = insert_table(doc, rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # 三线表边框设置
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        '  <w:top w:val="single" w:sz="12" w:space="0" w:color="000000"/>\n'
        '  <w:bottom w:val="single" w:sz="12" w:space="0" w:color="000000"/>\n'
        '  <w:left w:val="none"/>\n'
        '  <w:right w:val="none"/>\n'
        '  <w:insideH w:val="none"/>\n'
        '  <w:insideV w:val="none"/>\n'
        '</w:tblBorders>'
    )
    tblPr.append(tblBorders)

    # 表头行设置 (跨页重复表头与防止拆行)
    hdr_row = table.rows[0]
    hdr_trPr = hdr_row._tr.get_or_add_trPr()
    hdr_trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    hdr_trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

    hdr_cells = hdr_row.cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].text = ''
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_p_spacing(p, before=30, after=30)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(h_text)
        set_run_font(r, east_asia='SimSun', size_pt=9.5, bold=True)
        # 表头下细线
        tcPr = hdr_cells[i]._tc.get_or_add_tcPr()
        tcBorders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>\n'
            '  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
            '</w:tcBorders>'
        )
        tcPr.append(tcBorders)

    # 数据行
    for r_idx, r_data in enumerate(rows_data):
        row = table.rows[r_idx + 1]
        row_trPr = row._tr.get_or_add_trPr()
        row_trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        row_cells = row.cells
        for c_idx, cell_value in enumerate(r_data):
            row_cells[c_idx].text = ''
            p = row_cells[c_idx].paragraphs[0]
            if len(str(cell_value)) > 20:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_p_spacing(p, before=24, after=24)
            r = p.add_run(str(cell_value))
            set_run_font(r, east_asia='SimSun', size_pt=9.0)

    # 设置列宽
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)

    # 表格下方微小空行
    p_sp = insert_p(doc)
    set_p_spacing(p_sp, before=0, after=50)

print("Core builder with anchor support ready.")
