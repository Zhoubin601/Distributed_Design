"""Verify the reference-only v4 revision and render visual evidence."""
from pathlib import Path
from hashlib import sha256
from collections import Counter
import json
import re

import fitz
from PIL import Image, ImageDraw
from docx import Document

ROOT=Path(__file__).resolve().parents[2]
WORK=ROOT/'docs/20261024-reference-audit-v1'
DOCX=WORK/'20261024-misttable-distributed-design-v4.docx'
PDF=DOCX.with_suffix('.pdf')
OLD=ROOT/'output/v3/20261024-misttable-distributed-design-v3.docx'

def hash_file(p):
    return sha256(p.read_bytes()).hexdigest()

def values(t):
    return [[''.join(c.xpath('.//w:t/text()')) for c in row._tr.tc_lst] for row in t.rows]

def headings(d):
    return [(p.style.name,p.text) for p in d.paragraphs if p.style.name.lower().startswith('heading')]

old=Document(OLD)
new=Document(DOCX)
manifest=json.loads((WORK/'change-manifest.json').read_text('utf-8'))
checks={
    'v3_source_unchanged':hash_file(OLD)==manifest['source_sha256'],
    'headings_unchanged':headings(old)==headings(new),
    'paragraph_count_unchanged':len(old.paragraphs)==len(new.paragraphs),
    'table_count_unchanged':len(old.tables)==len(new.tables),
    'image_count_unchanged':len(old.inline_shapes)==len(new.inline_shapes),
    'cover_team_values_unchanged':values(old.tables[0])==values(new.tables[0]),
    'appendix_values_unchanged':values(old.tables[-1])==values(new.tables[-1]),
    'no_placeholders':not re.search('待填写|待确认|TODO','\n'.join(p.text for p in new.paragraphs)),
    'all_14_reviewed_changes_present':all(
        any(c['after']==p.text for p in new.paragraphs) or
        any(c['after']==p.text for t in new.tables for row in t.rows for cell in row.cells for p in cell.paragraphs)
        for c in manifest['changes']),
}
refs=[p.text for p in new.paragraphs if re.match(r'^\[\d+\]',p.text)]
checks['15_reference_entries']=len(refs)==15
checks['reference_fonts_consistent_9pt']=all(
    r.font.size is not None and r.font.size.pt==9
    for p in new.paragraphs if re.match(r'^\[\d+\]',p.text)
    for r in p.runs if r.text.strip())
changed_body=[c['after'] for c in manifest['changes'] if not re.match(r'^\[\d+\]',c['after'])]
checks['revised_body_citations_remain_superscript']=all(
    r.font.superscript is True for p in new.paragraphs if p.text in changed_body
    for r in p.runs if re.search(r'\[\d+(?:,\s*\d+)*\]',r.text))
checks['reference_9_real_access_date']='[2026-10-24]' in next(p for p in refs if p.startswith('[9] '))
checks['reference_5_pages_correct']='9-17' in next(p for p in refs if p.startswith('[5] '))
checks['reference_13_proceedings_correct']='Stabilization, Safety, and Security of Distributed Systems' in next(p for p in refs if p.startswith('[13] '))
checks['incorrect_forward_recovery_term_removed']=not any('用于前向恢复' in p.text for p in new.paragraphs)
pdf=fitz.open(PDF)
checks['pdf_26_pages']=len(pdf)==26
checks['all_pdf_pages_nonempty']=all(len(p.get_text().strip())>20 for p in pdf)
checks['pdf_access_date_visible']='2026-10-24' in pdf[24].get_text()
checks['pdf_backward_recovery_visible']='后向恢复' in pdf[22].get_text()
font_counts=Counter(span['font']+':'+str(round(span['size'],1))
    for i in range(3,25) for block in pdf[i].get_text('dict')['blocks'] if block['type']==0
    for line in block['lines'] for span in line['spans'])
for start in range(0,len(pdf),9):
    sheet=Image.new('RGB',(1500,2220),'#d8d8d8')
    draw=ImageDraw.Draw(sheet)
    for k,i in enumerate(range(start,min(start+9,len(pdf)))):
        pix=pdf[i].get_pixmap(matrix=fitz.Matrix(0.8,0.8),alpha=False)
        thumb=Image.frombytes('RGB',(pix.width,pix.height),pix.samples)
        thumb.thumbnail((480,700))
        x,y=(k%3)*500+10,(k//3)*740+30
        sheet.paste(thumb,(x,y));draw.text((x,y-20),f'Page {i+1}',fill='black')
    sheet.save(WORK/f'contact-{start//9+1}.png')
for i in (7,8,11,13,17,22,23,25,26):
    pdf[i-1].get_pixmap(matrix=fitz.Matrix(1.5,1.5),alpha=False).save(WORK/f'page-{i:02}.png')
summary={'checks':checks,'failed':[k for k,v in checks.items() if not v],
    'pages':len(pdf),'font_counts':dict(font_counts),'reference_entries':refs,
    'files':{p.name:{'sha256':hash_file(p),'bytes':p.stat().st_size} for p in (DOCX,PDF)}}
(WORK/'final-verification.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'checks':checks,'failed':summary['failed'],'pages':len(pdf)},ensure_ascii=False,indent=2))
if summary['failed']:
    raise SystemExit(1)
