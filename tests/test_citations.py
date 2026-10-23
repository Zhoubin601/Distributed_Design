# -*- coding: utf-8 -*-
"""
自动化测试套件：文献引用生成与格式验证
覆盖任务一要求的六大测试场景：
1. 单个文献引用
2. 多个文献的组合引用
3. 原始正文已经包含引用编号的情况（避免[9][9]等重复编号）
4. 连续出现多个引用锚点的段落
5. 最终生成的引用编号与文末参考文献编号完全一致闭包（1~15）
6. 引用编号确实以预期的上标格式及<w:noProof/>属性输出
"""

import sys
import os
import re
import pytest
import docx
from docx.oxml.ns import qn

# 确保脚本路径可导入
sys.path.insert(0, os.path.abspath('docs/scripts'))
import scratch_report_builder_core as core
import scratch_build_ch8_ch9_refs as refs_builder

def test_single_citation():
    """测试场景1：单个文献引用"""
    text = "分布式系统架构参考标准模型[1]，其设计原则贯穿始终。"
    segments = core.parse_body_text_segments(text)
    
    # 期望切分为3段：文字、上标[1]、文字
    assert len(segments) == 3
    assert segments[0] == ("分布式系统架构参考标准模型", False)
    assert segments[1] == ("[1]", True)
    assert segments[2] == ("，其设计原则贯穿始终。", False)

def test_multiple_citations_combo():
    """测试场景2：多个文献的组合引用"""
    text = "状态机复制与拜占庭容错机制分别参考经典理论[1, 7]，并在分布式环境中实施验证。"
    segments = core.parse_body_text_segments(text)
    
    assert len(segments) == 3
    assert segments[0] == ("状态机复制与拜占庭容错机制分别参考经典理论", False)
    assert segments[1] == ("[1, 7]", True)
    assert segments[2] == ("，并在分布式环境中实施验证。", False)
    
    # 测试带空格的情况
    text_space = "对比参考[1,  10]的系统实现"
    segments_space = core.parse_body_text_segments(text_space)
    assert segments_space[1] == ("[1,  10]", True)

def test_existing_citations_in_text_no_duplicates():
    """测试场景3：原始正文已包含引用编号，避免重复追加"""
    # 场景A：自然文本已写[9]，无需外部citations映射
    text_a = "参考优秀的轻量化虚拟桌游实践（如 Owlbear Rodeo [9]）进行房间生命周期管理"
    doc_a = docx.Document()
    p_a = core.add_body_p(doc_a, text_a)
    full_text_a = "".join(r.text for r in p_a.runs)
    assert "[9][9]" not in full_text_a
    assert full_text_a.count("[9]") == 1
    
    # 场景B：向后兼容传入citations=[('Owlbear Rodeo [9]', '[9]')]，也不得产生重复
    doc_b = docx.Document()
    p_b = core.add_body_p(doc_b, text_a, citations=[('Owlbear Rodeo [9]', '[9]')])
    full_text_b = "".join(r.text for r in p_b.runs)
    assert "[9][9]" not in full_text_b
    assert full_text_b.count("[9]") == 1
    
    # 场景C：正文已有[1, 7]，传入[7]不得被截断或重复
    text_c = "参考经典共识与复制模型[1, 7]进行权衡"
    doc_c = docx.Document()
    p_c = core.add_body_p(doc_c, text_c, citations=[('模型[1, 7]', '[7]')])
    full_text_c = "".join(r.text for r in p_c.runs)
    assert "[1, 7][7]" not in full_text_c
    assert full_text_c.count("[1, 7]") == 1

def test_consecutive_multiple_citations_in_paragraph():
    """测试场景4：连续出现多个引用锚点的段落"""
    text = "分布式系统架构遵从[1]，采用两阶段提交[3]实现事务一致性，并通过心跳保活机制[12]感知故障。"
    segments = core.parse_body_text_segments(text)
    
    # 统计上标数量与顺序
    sup_segments = [seg for seg, is_sup in segments if is_sup]
    assert sup_segments == ["[1]", "[3]", "[12]"]
    
    doc = docx.Document()
    p = core.add_body_p(doc, text)
    sup_runs = [r.text for r in p.runs if r.font.superscript]
    assert sup_runs == ["[1]", "[3]", "[12]"]

def test_citations_correspondence_with_references():
    """测试场景5：最终正文中所有引用的编号与文末参考文献编号一致且完备闭包"""
    # 获取参考文献库中的编号列表
    ref_keys = [re.match(r'\[(\d+)\]', r).group(1) for r in refs_builder.REFERENCES]
    assert len(ref_keys) == 15
    assert ref_keys == [str(i) for i in range(1, 16)]
    
    # 检查现有的实际交付文档中的所有引用编号
    doc_path = "output/Distributed_Design_Report_Draft_v2.docx" if os.path.exists("output/Distributed_Design_Report_Draft_v2.docx") else "output/Distributed_Design_Report_Draft_v1.docx"
    assert os.path.exists(doc_path), f"交付文档不存在: {doc_path}"
    
    doc = docx.Document(doc_path)
    all_cited_ids = set()
    
    # 提取所有段落与表格中的引用编号（合并连续上标run以兼容Word COM词法分词）
    for p in doc.paragraphs:
        # 跳过文末参考文献列表本身所在的段落
        if "【参考文献】" in p.text or (p.text.startswith("[") and "]" in p.text and any(f"[{i}]" in p.text for i in range(1, 16)) and len(p.text) > 30):
            continue
        curr_sup = []
        for r in p.runs:
            if r.font.superscript:
                curr_sup.append(r.text)
            else:
                if curr_sup:
                    sup_text = ''.join(curr_sup)
                    for group in re.findall(r'\[(\d+(?:,\s*\d+)*)\]', sup_text):
                        for num in re.split(r',\s*', group):
                            all_cited_ids.add(num.strip())
                    curr_sup = []
        if curr_sup:
            sup_text = ''.join(curr_sup)
            for group in re.findall(r'\[(\d+(?:,\s*\d+)*)\]', sup_text):
                for num in re.split(r',\s*', group):
                    all_cited_ids.add(num.strip())

    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    curr_sup = []
                    for r in p.runs:
                        if r.font.superscript:
                            curr_sup.append(r.text)
                        else:
                            if curr_sup:
                                sup_text = ''.join(curr_sup)
                                for group in re.findall(r'\[(\d+(?:,\s*\d+)*)\]', sup_text):
                                    for num in re.split(r',\s*', group):
                                        all_cited_ids.add(num.strip())
                                curr_sup = []
                    if curr_sup:
                        sup_text = ''.join(curr_sup)
                        for group in re.findall(r'\[(\d+(?:,\s*\d+)*)\]', sup_text):
                            for num in re.split(r',\s*', group):
                                all_cited_ids.add(num.strip())
                                    
    # 验证：引用的编号必须全部属于 1~15
    invalid_ids = all_cited_ids - set(ref_keys)
    assert len(invalid_ids) == 0, f"存在未在参考文献中定义的引用编号: {invalid_ids}"
    
    # 验证：1~15的参考文献必须全部被引用，无冗余未引用的参考文献条目
    uncited_ids = set(ref_keys) - all_cited_ids
    assert len(uncited_ids) == 0, f"存在未被正文引用的参考文献编号: {uncited_ids}"

def test_citation_superscript_format():
    """测试场景6：引用编号以预期的上标格式输出，且包含<w:noProof/>"""
    doc = docx.Document()
    text = "状态机复制机制[7]保证了强一致状态演进。"
    p = core.add_body_p(doc, text)
    
    found_sup_run = False
    for r in p.runs:
        if r.text == "[7]":
            found_sup_run = True
            # 验证 python-docx 属性
            assert r.font.superscript is True
            
            # 验证底层 XML 属性
            rPr = r._r.get_or_add_rPr()
            vertAlign = rPr.find(qn('w:vertAlign'))
            assert vertAlign is not None
            assert vertAlign.get(qn('w:val')) == 'superscript'
            
            # 验证防止拼写检查分割 run 的 noProof 标记
            noProof = rPr.find(qn('w:noProof'))
            assert noProof is not None
            
    assert found_sup_run, "未找到目标上标运行块 [7]"

if __name__ == "__main__":
    pytest.main(["-v", __file__])
