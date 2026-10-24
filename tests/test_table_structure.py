# -*- coding: utf-8 -*-
"""
自动化测试套件：表格结构属性与分页保护验证
覆盖任务一要求的四大测试场景以及整体验收断言：
1. 正常表格结构属性全部合规通过
2. 缺少 w:tblHeader 时独立报错失败 (且不跳过其他检查)
3. 缺少表头段落 keep_with_next 时独立报错失败
4. 缺少数据行 w:cantSplit 时独立报错失败
5. 验证实际报告中全部 10 张专业数据表结构合规性
6. 验证 ReportVerifier 针对损坏属性生成明确 FAIL 结果并拒绝通过
"""

import sys
import os
import pytest
import docx
from docx.oxml.ns import qn

sys.path.insert(0, os.path.abspath('docs/scripts'))
import scratch_report_builder_core as core
import scratch_verify_output as verifier_mod
from report_paths import resolve_report_docx, resolve_report_pdf

def build_dummy_table(rows_count=3, cols_count=3):
    """使用报告核心构建函数构造标准三线表测试样例"""
    doc = docx.Document()
    headers = [f"表头列 {i+1}" for i in range(cols_count)]
    rows_data = [[f"数据 ({r+1},{c+1})" for c in range(cols_count)] for r in range(rows_count)]
    core.add_table_with_caption(doc, "99-1", "测试验证表格", headers, rows_data)
    # 提取生成的表格
    tables = verifier_mod.get_content_tables(doc)
    assert len(tables) == 1
    return doc, tables[0][2]

def test_normal_table_passes():
    """测试场景 1：标准生成的表格应完整具备三大分页保护属性并全部通过"""
    _, tbl = build_dummy_table()
    res = verifier_mod.verify_single_table_structure(tbl, "99-1")
    
    assert res["has_tbl_header"] is True, "正常表头行必须具备 w:tblHeader"
    assert res["all_header_p_keep"] is True, "正常表头段落必须全部设置 keep_with_next"
    assert res["all_rows_cant_split"] is True, "正常表格所有行必须全部设置 w:cantSplit"
    assert res["is_valid"] is True, "整体结构验证必须通过"
    assert len(res["missing_header_p_cells"]) == 0
    assert len(res["missing_cant_split_rows"]) == 0

def test_missing_tbl_header_fails():
    """测试场景 2：表头行缺少 w:tblHeader 时失败，且不能跳过后续 keep_with_next 检查"""
    _, tbl = build_dummy_table()
    
    # 移除表头行的 w:tblHeader
    hdr_trPr = tbl.rows[0]._tr.trPr
    assert hdr_trPr is not None
    tbl_hdr = hdr_trPr.find(qn('w:tblHeader'))
    assert tbl_hdr is not None
    hdr_trPr.remove(tbl_hdr)
    
    res = verifier_mod.verify_single_table_structure(tbl, "99-1")
    assert res["has_tbl_header"] is False, "缺少 w:tblHeader 必须被准确检测"
    assert res["is_valid"] is False, "缺少 w:tblHeader 必须判定整个表格无效"
    # 关键断言：即使缺少 w:tblHeader，keep_with_next 检查也不能被跳过
    assert res["all_header_p_keep"] is True, "未破坏 keep_with_next 时该项应被独立检出并通过"
    assert res["all_rows_cant_split"] is True, "未破坏 cantSplit 时该项应被独立检出并通过"

def test_missing_keep_with_next_fails():
    """测试场景 3：表头单元格缺少 keep_with_next 时失败"""
    _, tbl = build_dummy_table()
    
    # 破坏第 1 个表头单元格的 keep_with_next
    target_p = tbl.rows[0].cells[0].paragraphs[0]
    target_p.paragraph_format.keep_with_next = False
    pPr = target_p._p.find(qn('w:pPr'))
    if pPr is not None:
        kn = pPr.find(qn('w:keepNext'))
        if kn is not None:
            pPr.remove(kn)
            
    res = verifier_mod.verify_single_table_structure(tbl, "99-1")
    assert res["has_tbl_header"] is True, "w:tblHeader 保持存在"
    assert res["all_header_p_keep"] is False, "表头缺少 keep_with_next 必须被检出"
    assert (0, 0) in res["missing_header_p_cells"]
    assert res["is_valid"] is False, "结构必须判定无效"

def test_missing_cant_split_fails():
    """测试场景 4：数据行缺少 w:cantSplit 时失败"""
    _, tbl = build_dummy_table(rows_count=4)
    
    # 破坏第 2 行数据行 (index 2) 的 w:cantSplit
    r2_trPr = tbl.rows[2]._tr.trPr
    assert r2_trPr is not None
    cs = r2_trPr.find(qn('w:cantSplit'))
    assert cs is not None
    r2_trPr.remove(cs)
    
    res = verifier_mod.verify_single_table_structure(tbl, "99-1")
    assert res["has_tbl_header"] is True
    assert res["all_header_p_keep"] is True
    assert res["all_rows_cant_split"] is False, "数据行缺少 cantSplit 必须被检出"
    assert 2 in res["missing_cant_split_rows"]
    assert res["is_valid"] is False, "结构必须判定无效"

def test_real_doc_content_tables_all_valid():
    """测试场景 5：验证实际报告文件 output/Distributed_Design_Report_Draft_v2.docx 中全部 10 张专业数据表结构"""
    docx_path = resolve_report_docx()
    assert os.path.exists(docx_path), f"报告文件必须存在: {docx_path}"
    doc = docx.Document(docx_path)
    
    content_tables = verifier_mod.get_content_tables(doc)
    assert len(content_tables) == 10, f"必须准确识别正文 10 个专业数据表，实测: {len(content_tables)}"
    
    expected_ids = ["1-1", "1-2", "2-1", "3-1", "4-1", "5-1", "6-1", "7-1", "8-1", "9-1"]
    found_ids = [t[0] for t in content_tables]
    assert found_ids == expected_ids, f"识别到的表格编号与顺序不符: {found_ids}"
    
    for tid, cap, tbl in content_tables:
        res = verifier_mod.verify_single_table_structure(tbl, tid)
        assert res["has_tbl_header"] is True, f"表 {tid} ({cap}) 缺失 w:tblHeader 属性"
        assert res["all_header_p_keep"] is True, f"表 {tid} ({cap}) 表头单元格存在未启用 keep_with_next 的段落: {res['missing_header_p_cells']}"
        assert res["all_rows_cant_split"] is True, f"表 {tid} ({cap}) 存在未启用 cantSplit 的数据行: {res['missing_cant_split_rows']}"
        assert res["is_valid"] is True, f"表 {tid} ({cap}) 结构验证未完全通过"

def test_verifier_produces_fail_on_tampered_doc():
    """测试场景 6：验证 ReportVerifier 在属性缺失时确实产生明确 FAIL 结果，拒绝通过"""
    docx_path = resolve_report_docx()
    doc = docx.Document(docx_path)
    content_tables = verifier_mod.get_content_tables(doc)
    assert len(content_tables) >= 1
    
    # 人为篡改表 1-1 的表头，移除 w:tblHeader
    t1_1 = content_tables[0][2]
    hdr_trPr = t1_1.rows[0]._tr.trPr
    tbl_hdr = hdr_trPr.find(qn('w:tblHeader'))
    assert tbl_hdr is not None
    hdr_trPr.remove(tbl_hdr)
    
    verifier = verifier_mod.ReportVerifier(docx_path, resolve_report_pdf())
    verifier.verify_figures_and_tables(doc)
    
    # 检查是否记录了 FAIL
    fail_items = [r for r in verifier.results if r[2] == "FAIL"]
    assert len(fail_items) > 0, "篡改后必须产生至少一项 FAIL 记录"
    tbl_hdr_fail = any("w:tblHeader" in r[1] for r in fail_items)
    assert tbl_hdr_fail is True, "必须明确指出缺失 w:tblHeader"
